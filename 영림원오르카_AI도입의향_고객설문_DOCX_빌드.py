#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
「영림원소프트랩·오르카아이티 AI 도입 의향 조사」 설문 설계안 -> DOCX 생성.

배경: 정춘용 CMO 지시(2026-08-05). 영림원소프트랩·오르카아이티 기존 ERP 고객사를 대상으로
AI 도입의지·투자의지·정부 지원사업(스마트공장구축 - 제조AI 특화) 참여가능성을 확인하는 설문을
양사 공식 레터 형태로 발송하기 위한 초안. 사용자가 제시한 7개 원안 문항에 답변항목을 붙이고,
리드 확인(응답자 자격·예산·경쟁상황·과거경험·후속미팅) 문항을 보강했다.

이 설문은 기존 전략기획 문서(`영림원AI연합군_사업전략_병합_PPTX_빌드.py` NEW-B 슬라이드,
"2026-09 공식레터 설문 -> 방문 -> 사업계획서 작성(~12월) -> 2027-01 사업신청")에서 이미
구상되어 있던 실행절차의 1단계를 구체화한 것이다.

사실관계 주의(중요): 정부지원사업(스마트공장구축 - 제조AI 특화, 구 표현 "AI-Agent 지원사업")은
2026-08 현재 2027년도 공고가 아직 발표되지 않았다(연초 통합공고 패턴, 출처:
`05_공모전/00_공통_지원자료/공고스캔_S급품목_2026-08-04.md`). "총사업비 50% · 최대 2억원"은
2026년 선정 4건(저스템·알피에스·선영코리아·스피폭스)을 관찰한 결과이지 정부 공식 규정이 아니다
(출처: `01_마케팅/캠페인/2026-07_영림원골드파트너_AIFactory_공동마케팅/분석_영림원기존고객사_이익분석.md`
§2.3, `기획안_2027제조AI특화제안.md`). 고객 발송용 문서이므로 본문에 "관찰 기준" 라벨을 명시했다.

실행: python 영림원오르카_AI도입의향_고객설문_DOCX_빌드.py
출력: 영림원오르카_AI도입의향_고객설문_설계안_2026-08-05.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "영림원오르카_AI도입의향_고객설문_설계안_2026-08-05.docx")

NAVY = RGBColor(0x1E, 0x27, 0x61)
CORAL = RGBColor(0xB0, 0x30, 0x30)
GRAY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x14, 0x6B, 0x44)

doc = Document()

# ---- 기본 스타일 ----
style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10.5)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find(qn("w:rFonts"))
if rFonts is None:
    rFonts = OxmlElement("w:rFonts")
    rpr.append(rFonts)
rFonts.set(qn("w:eastAsia"), "맑은 고딕")

for sec in doc.sections:
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)


def h1(text, color=NAVY):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = color
    r.font.size = Pt(17)
    r.font.name = "맑은 고딕"
    return p


def h2(text, color=NAVY):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = color
    r.font.size = Pt(13.5)
    r.font.name = "맑은 고딕"
    return p


def h3(text, color=NAVY):
    p = doc.add_heading(level=3)
    r = p.add_run(text)
    r.font.color.rgb = color
    r.font.size = Pt(11.5)
    r.font.name = "맑은 고딕"
    return p


def para(text, size=10.5, bold=False, italic=False, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run("※ " + text)
    r.font.size = Pt(9)
    r.italic = True
    r.font.color.rgb = GRAY
    return p


def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def question_block(qnum, qtext, options, qtype="단일 선택", subnote=None, badge=None, badge_color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{qnum}. {qtext}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    if badge:
        rb = p.add_run(f"   [{badge}]")
        rb.font.size = Pt(8.5)
        rb.bold = True
        rb.font.color.rgb = badge_color
    pt = doc.add_paragraph()
    pt.paragraph_format.space_after = Pt(4)
    pt.paragraph_format.left_indent = Cm(0.3)
    rt = pt.add_run(f"({qtype})")
    rt.font.size = Pt(9)
    rt.italic = True
    rt.font.color.rgb = GRAY
    for opt in options:
        po = doc.add_paragraph()
        po.paragraph_format.left_indent = Cm(0.6)
        po.paragraph_format.space_after = Pt(2)
        ro = po.add_run(opt)
        ro.font.size = Pt(10.5)
    if subnote:
        note(subnote)


def review_row(table, qlabel, comment, addition=None):
    row = table.add_row()
    row.cells[0].text = qlabel
    row.cells[1].text = comment
    row.cells[2].text = addition or "-"
    for c in row.cells:
        for pgh in c.paragraphs:
            for rn in pgh.runs:
                rn.font.size = Pt(9.5)


# ============================================================ 표지
h1("영림원소프트랩 · 오르카아이티 AI 도입 의향 조사")
para("고객 대상 설문 설계안 (초안) — 정부 스마트공장 제조AI 특화 지원사업 연계", size=12.5, bold=True, color=GRAY)
para("작성: 워크벤치(클로드) 초안 · 검토·확정: 정춘용 CMO   |   작성일: 2026-08-05", size=9.5, color=GRAY)
hr()

box = doc.add_paragraph()
box.paragraph_format.space_after = Pt(4)
rb = box.add_run(
    "이 문서 사용법 — 이 파일은 (1) 고객에게 바로 발송 가능한 형태로 다듬은 \"공식 서한 + 설문지\"(PART 2~5)와, "
    "(2) 정춘용 CMO의 검토를 위한 \"워크벤치 검토의견\"(PART 6, 문항별 코멘트·추가 제안·사실관계 주의사항)으로 구성됩니다. "
    "PART 6은 내부 검토용이므로 실제 고객 발송본에서는 반드시 제외하십시오."
)
rb.font.size = Pt(9.5)
rb.italic = True
rb.font.color.rgb = CORAL

doc.add_page_break()

# ============================================================ PART 1
h2("PART 1. 설문 설계 배경 (요약)")

para("목적", bold=True, size=11)
para(
    "영림원소프트랩·오르카아이티의 기존 ERP 고객사를 대상으로 AI 도입에 대한 관심도·투자의지·정부 지원사업 "
    "참여가능성을 파악하여, AI 캠페인 전략 수립과 정부 지원사업(스마트공장구축 - 제조AI 특화) 신청 후보군 "
    "선별의 근거 데이터를 확보한다."
)

para("출발 가설 (정춘용 CMO 구술, 2026-08-05)", bold=True, size=11)
para(
    "2026년 8월 현재부터 2028년 연말까지 약 2년 5개월 동안, 영림원 K-System·QAD ERP를 사용하는 IT 수준의 "
    "중소·중견기업 대부분이 AI 관련 프로젝트를 검토하고 실제 예산을 편성해 실행할 것이라는 전망에서 출발한다. "
    "이 설문은 그 전망이 실제 고객 응답으로 뒷받침되는지 확인하고, 동시에 정부 지원사업을 활용하도록 유도하는 "
    "1차 접점 역할을 한다."
)

para("대상 모집단 (참고 — 영림원 쪽 기존 관찰치, 오르카아이티는 별도 확인 필요)", bold=True, size=11)
para(
    "영림원 K-System Ace 기준 국내외 약 2,600여 고객사 중, 고객가치사업부 직할고객 약 800개, 그 중 제조업이 "
    "약 600개(75%) 수준으로 추정된다(정춘용 CMO 판단치, 공식 통계 아님 — 출처: "
    "01_마케팅/캠페인/2026-07_영림원골드파트너_AIFactory_공동마케팅/기획안_영림원AI협력전략기획서.md). "
    "정부 스마트공장 지원사업은 제조업을 대상으로 하므로, 비제조 고객군은 이번 설문의 1차 우선순위에서 "
    "제외하거나 별도 문항으로 구분하는 것을 권장한다. 오르카아이티(QAD ERP) 고객 모집단 규모는 이 워크스페이스에 "
    "확인된 자료가 없어 별도 확인이 필요하다."
)

para("실행 절차와의 연계", bold=True, size=11)
para(
    "이 설문은 기존 사업전략 기획(영림원AI연합군_사업전략_병합_PPTX_빌드.py, NEW-B \"실행 절차\" 슬라이드)에서 "
    "이미 구상된 흐름의 1단계다: 공식레터 설문 → 고객 방문·요구사항 파악 → AI연합군 준비사항 설명 → "
    "AI프로젝트 범위·목표 설정 → 사업계획서 작성 → 정부 지원사업 신청. 해당 기획서는 설문을 \"2026-09\" 시행, "
    "신청을 \"2027-01\"로 가정했으나 이는 가상 시나리오였다 — 이번 설문 설계로 첫 단계가 구체화된다."
)

para("중요 사실관계 — 정부 지원사업 표기 원칙 (반드시 확인)", bold=True, size=11, color=CORAL)
p = doc.add_paragraph()
r = p.add_run(
    "\"AI-Agent 지원사업\"은 WACE·영림원 캠페인 내부에서 쓰는 마케팅상 별칭이며, 정부의 공식 사업명은 "
    "\"스마트공장구축 - 제조AI 특화\"(자율형공장 AI트랙)다. 2026-08-05 현재 2027년도 공고는 아직 발표되지 "
    "않았다(연초 통합공고 패턴 — 출처: 05_공모전/00_공통_지원자료/공고스캔_S급품목_2026-08-04.md). "
    "\"현재 진행 중\"이 아니라 \"매년 초 공고되는 사업이며 2027년분은 공고 예정\"으로 표기해야 한다. "
    "또한 \"총사업비의 50%, 최대 2억원\"은 정부의 공식 지원 규정을 직접 인용한 것이 아니라 2026년 선정 4건"
    "(저스템·알피에스·선영코리아·스피폭스)을 관찰한 결과다(출처: 분석_영림원기존고객사_이익분석.md §2.3, "
    "기획안_2027제조AI특화제안.md). 고객에게 발송되는 공식 문서이므로 아래 서한·설문 문항에는 "
    "\"2026년 선정 사례 기준\"이라는 관찰 기준 라벨을 넣었다 — 문구를 임의로 단순화(\"정부가 50%를 지원합니다\" "
    "식 확정 표현)하지 않도록 주의를 권한다."
)
r.font.size = Pt(9.8)

doc.add_page_break()

# ============================================================ PART 2 — 공식 서한
h2("PART 2. 공식 서한 (설문지 표지 겸 커버레터)")
note("발송 시 아래 [ ] 표시 부분을 발신 회사에 맞춰 확정하고, 두 회사 명의로 각각 1부씩(영림원소프트랩용 / 오르카아이티용) 별도로 완성해 발송할 것을 권장합니다. 담당자·직인·로고는 각 사에서 채워 넣습니다.")

letter = doc.add_paragraph()
letter.paragraph_format.space_after = Pt(4)
letter.add_run("[ 영림원소프트랩 / 오르카아이티 ]  AI사업부").bold = True

para("고객 AI 도입 의향 조사 안내", bold=True, size=13, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

body_lines = [
    "안녕하십니까. 귀사의 지속적인 성장과 발전을 기원합니다.",
    "",
    "[ 영림원소프트랩 / 오르카아이티 ]는 귀사가 오랫동안 함께해 주신 ERP 파트너로서, 최근 제조·경영 현장 "
    "전반에 확산되고 있는 AI(인공지능) 도입 흐름에 대해 귀사의 의견을 여쭙고자 이 서한을 보내드립니다.",
    "",
    "저희는 향후 2~3년 내 대부분의 제조·중견기업이 데이터 기반 AI 프로젝트를 검토하고 실제 투자를 집행하게 "
    "될 것으로 전망하고 있습니다. 특히 귀사와 같이 이미 ERP 시스템을 통해 축적된 데이터를 보유한 기업은 "
    "그렇지 않은 기업보다 AI 도입에 유리한 조건을 갖추고 있습니다.",
    "",
    "이와 관련해 정부는 매년 초 \"스마트공장구축 - 제조AI 특화\"(통칭 AI-Agent 지원사업) 공고를 통해 "
    "AI 시스템을 도입하려는 기업에 사업비 일부를 지원하고 있습니다. 2026년 선정 사례 기준으로는 총사업비 "
    "약 4억원 중 정부지원 최대 2억원 수준(자기부담 2억원)으로 진행된 바 있습니다(2027년도 공고는 아직 "
    "발표 전이며, 세부 조건은 공고 시 확정됩니다).",
    "",
    "이번 설문은 귀사의 AI 도입 관심도와 계획을 파악하고, 필요하신 경우 저희 AI사업부가 정부 지원사업 신청 "
    "과정을 처음부터 함께 안내해 드리기 위한 목적으로 진행합니다. 응답에는 약 5분이 소요되며, 수집된 정보는 "
    "본 조사 목적 외에는 사용하지 않습니다.",
    "",
    "바쁘신 중에도 아래 설문에 참여해 주시면 귀사의 AI 도입 계획을 구체화하는 데 실질적인 도움을 드릴 수 "
    "있도록 최선을 다하겠습니다. 감사합니다.",
]
for line in body_lines:
    if line == "":
        doc.add_paragraph()
    else:
        para(line, size=10.5, space_after=4)

para("[ 영림원소프트랩 / 오르카아이티 ] AI사업부장  [ 성명 ]  배상", size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
para("문의: [ 담당자 성명 / 전화 / 이메일 ]", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=16)

doc.add_page_break()

# ============================================================ PART 3 — 응답자 기본정보
h2("PART 3. 응답자 기본정보")
basic_items = [
    "A. 회사명 : ______________________________",
    "B. 담당자 성명 / 직책 : ______________________________ / ______________________________",
    "C. 연락처 (전화 / 이메일) : ______________________________ / ______________________________",
    "D. 업종 : ☐ 제조업(세부: __________)  ☐ 유통/서비스  ☐ 건설  ☐ 기타(__________)",
    "E. 종업원 수 : ☐ 10인 미만  ☐ 10~49인  ☐ 50~99인  ☐ 100~299인  ☐ 300인 이상",
    "F. 연매출 규모(대략) : ☐ 50억 미만  ☐ 50~100억  ☐ 100~300억  ☐ 300~500억  ☐ 500억 이상",
    "G. 현재 사용 중인 ERP 및 도입연차 : ☐ 영림원 K-System  ☐ QAD ERP  ☐ 기타(__________)  /  도입 후 약 ____년",
]
for it in basic_items:
    para(it, size=10.5, space_after=5)

doc.add_page_break()

# ============================================================ PART 4 — 본 설문
h2("PART 4. 본 설문 — AI 도입의지 · 투자의지 · 지원사업 참여가능성")
note("문항 옆 [신규] 표시는 사용자 원안 7문항을 보강하기 위해 워크벤치가 추가 제안한 문항입니다. 채택 여부는 PART 6 검토의견을 참고해 결정해 주십시오.")

question_block(
    "Q1", "귀사는 ERP시스템을 어느 정도 잘 활용하고 계십니까?",
    [
        "① 전사 핵심업무(생산·재고·회계 등) 전반에 걸쳐 적극적으로 활용하고 있다",
        "② 주요 업무 위주로 활용하고 있으나, 일부 기능은 제대로 활용하지 못하고 있다",
        "③ 기본적인 데이터 입력·조회 수준에 머물러 있다",
        "④ 도입은 했으나 전반적으로 활용도가 낮은 편이다",
        "⑤ 잘 모르겠다 / 담당 업무가 아니다",
    ],
    subnote="원안 1번. ERP 활용 성숙도를 확인하는 문항 — AI 적용 가능성은 ERP 데이터 축적·활용 수준과 비례하므로, 뒤 문항(Q3 이하)의 해석 기준이 됩니다.",
)

question_block(
    "Q2", "귀사는 ERP 외에 AI시스템 도입에 관심이 있으십니까?",
    [
        "① 매우 관심이 많다 (구체적으로 검토·조사 중이다)",
        "② 관심이 있다 (아직 구체적인 계획은 없다)",
        "③ 보통이다 (필요성은 느끼나 우선순위가 높지 않다)",
        "④ 관심이 적은 편이다",
        "⑤ 전혀 관심이 없다",
    ],
    subnote="원안 2번.",
)

question_block(
    "Q2-1", "귀하는 귀사의 AI 도입 관련 의사결정에 어느 정도 관여하고 계십니까?",
    [
        "① 최종 의사결정권자(대표/임원)이다",
        "② 의사결정에 실무적으로 깊이 관여한다 (예산·기획 담당)",
        "③ 관련 부서 실무자로서 의견을 제시하는 수준이다",
        "④ 의사결정과는 거리가 있다",
    ],
    badge="신규 — 응답 신뢰도 확인",
    badge_color=GREEN,
    subnote="응답자가 실제 예산·투자 결정에 영향을 미치는 위치인지 확인하는 문항입니다. 이 문항의 응답에 따라 이후 응답(특히 Q5-1, Q7)의 신뢰도 가중치를 다르게 부여할 수 있습니다.",
)

question_block(
    "Q3", "AI시스템을 실제 예산을 들여 도입한다면, 그 시기는 언제로 계획하십니까?",
    [
        "① 이미 도입했거나 진행 중이다",
        "② 2026년 내 (올해)",
        "③ 2027년 중",
        "④ 2028년 중",
        "⑤ 2029년 이후 또는 아직 계획이 없다",
        "⑥ 아직 시기를 정하지 못했다",
    ],
    subnote="원안 3번. 2028년 말까지의 전망을 검증하는 핵심 문항이므로, 연도별 구간을 세분화했습니다(원안은 단일 개방형 질문이었음).",
)

question_block(
    "Q4", "AI시스템을 도입한다면 어느 분야에 우선 도입하려 하십니까? (복수 응답 가능)",
    [
        "① 생산/제조 현장 (설비 예지보전 · 품질검사 · 공정 최적화 등)",
        "② 구매/자재/재고관리 (SCM)",
        "③ 영업/수요예측",
        "④ 회계/재무",
        "⑤ 인사/총무",
        "⑥ 고객상담/CS",
        "⑦ 경영진 의사결정 지원 (BI·경영 리포트)",
        "⑧ 기타 ( 직접 기재: ______________ )",
    ],
    qtype="복수 선택",
    subnote="원안 4번. 복수 응답으로 전환해 실제 우선순위 분포를 파악하도록 했습니다.",
)

question_block(
    "Q5", "AI 도입 이후 가장 기대하는 효과는 무엇입니까? (최대 2개 선택)",
    [
        "① 생산성 향상 / 인력 효율화",
        "② 불량률 감소 / 품질 향상",
        "③ 데이터 기반 의사결정 속도 향상",
        "④ 비용 절감",
        "⑤ 신규 매출 / 사업기회 창출",
        "⑥ 업무 표준화 및 프로세스 개선",
        "⑦ 기타 ( 직접 기재: ______________ )",
    ],
    qtype="복수 선택 (최대 2개)",
    subnote="원안 5번. \"최대 2개\"로 제한해 진짜 우선순위를 드러내도록 했습니다(전부 선택 시 변별력이 사라지는 것을 방지).",
)

question_block(
    "Q5-1", "AI 도입을 위한 예산 편성 현황은 어떻습니까?",
    [
        "① 이미 예산을 편성했거나 확보했다",
        "② 편성할 계획이 있으나 구체적 규모는 아직 정하지 못했다",
        "③ 예산 편성 여부를 검토 중이다",
        "④ 아직 예산 편성 계획이 없다",
    ],
    badge="신규 — 투자의지 핵심 문항",
    badge_color=GREEN,
    subnote="\"관심\"(Q2)과 \"실제 투자 실행\"은 다릅니다. 이 문항이 실질적인 투자의지를 가장 직접적으로 드러내는 질문입니다.",
)

question_block(
    "Q5-2", "현재 AI 도입을 위해 검토 중인 다른 업체(솔루션·벤더)가 있으십니까?",
    [
        "① 있다 — 구체적으로 진행 중인 업체가 있다",
        "② 있다 — 초기 검토·미팅 단계다",
        "③ 없다 — 아직 특정 업체를 정하지 않았다",
        "④ 없다 — 자체적으로 진행할 계획이다",
    ],
    badge="신규 — 경쟁상황 파악",
    badge_color=GREEN,
    subnote="①·② 응답 고객은 의사결정 속도가 빠른 대신 경쟁이 있는 고객이므로, 후속 컨택 우선순위를 판단하는 데 유용합니다.",
)

question_block(
    "Q6",
    "정부는 매년 초 \"스마트공장구축 - 제조AI 특화\"(통칭 AI-Agent 지원사업) 공고를 통해 AI 시스템을 도입하는 "
    "기업에 사업비 일부를 지원하고 있습니다. (2026년 선정 사례 기준: 총사업비 약 4억원 중 정부지원 최대 2억원 "
    "수준 — 2027년도 공고는 미발표이며 세부 조건은 공고 시 확정) 이러한 정부 지원사업에 참여하실 의향이 "
    "있으십니까?",
    [
        "① 매우 적극적으로 참여하고 싶다",
        "② 조건(지원 규모·자기부담 등)에 따라 참여를 검토하겠다",
        "③ 아직 판단하기 이르다 (추가 정보가 필요하다)",
        "④ 참여 의사가 없다",
    ],
    subnote="원안 6번을 사실관계에 맞게 수정했습니다. 원안의 \"전체 사업금액의 50%, 최대 2억원\"은 정부의 공식 지원 규정이 아니라 2026년 선정 4건을 관찰한 결과이므로, 확정 표현 대신 \"2026년 선정 사례 기준\"으로 표기했습니다. PART 1·PART 6 참고.",
)

question_block(
    "Q6-1", "과거 정부 지원사업(스마트공장 구축, 제조업 R&D 등)에 신청하거나 선정된 경험이 있으십니까?",
    [
        "① 있다 — 선정되어 사업을 수행한 경험이 있다",
        "② 있다 — 신청했으나 선정되지 못한 경험이 있다",
        "③ 없다 — 신청해본 적이 없다",
        "④ 없다 — 관련 정보를 잘 모른다",
    ],
    badge="신규 — 신청 경험 파악",
    badge_color=GREEN,
    subnote="경험이 있는 고객은 서류 준비 부담을 낮게 느끼고, 경험이 없는 고객은 \"AI사업부가 가이드\"(Q7)의 가치를 더 크게 느낄 가능성이 높습니다 — 두 그룹에 다른 접근이 필요합니다.",
)

question_block(
    "Q7",
    "저희 [ 영림원소프트랩 / 오르카아이티 ] AI사업부가 정부 지원사업 신청 준비부터 AI 프로젝트 수행까지 "
    "처음부터 함께 안내해 드린다면, 함께 신청하시겠습니까?",
    [
        "① 예, 적극적으로 함께 신청하고 싶다",
        "② 관심은 있으나 사전 설명(설명회·미팅)을 먼저 받아보고 싶다",
        "③ 아직 결정하기 이르다",
        "④ 신청 의사가 없다",
    ],
    subnote="원안 7번.",
)

question_block(
    "Q7-1", "(Q7에서 ① 또는 ②를 선택하신 경우) 담당자와의 후속 미팅을 희망하시는 시기를 알려주십시오.",
    [
        "① 가능한 빨리 (1~2주 이내)",
        "② 1개월 이내",
        "③ 2~3개월 이내",
        "④ 아직 시기를 정하기 어렵다",
    ],
    badge="신규 — 리드 전환 핵심 문항",
    badge_color=CORAL,
    subnote="이 설문의 최종 목적은 \"AI 도입의지가 확인된 고객이 실제로 우리와 함께 정부 지원사업을 신청할 의향이 있는지\"를 확인하는 것입니다. Q7-1이 그 의향을 실행(미팅 일정)으로 전환하는 마지막 문항이므로, 반드시 포함할 것을 권장합니다.",
)

doc.add_page_break()

# ============================================================ PART 5 — 마무리
h2("PART 5. 마무리 및 후속 연락 동의")
para(
    "개인정보 수집·이용 동의: 본 설문에서 수집하는 개인정보(담당자 성명·연락처)는 AI 도입 상담 및 정부 지원사업 "
    "안내 목적으로만 사용되며, 목적 달성 후 관련 법령에 따라 파기됩니다. 동의를 거부하실 수 있으며, 이 경우 "
    "후속 상담 안내가 제한될 수 있습니다.",
    size=9.8,
)
para("☐ 위 내용에 동의하며, 설문에 성실히 응답합니다.", size=10.5, space_after=14)
para("담당자 서명: ______________________   날짜: 2026년 ____월 ____일", size=10.5, space_after=4)
para("설문에 참여해 주셔서 진심으로 감사드립니다. 회신 주신 내용은 담당자가 개별 확인 후 순차적으로 연락드리겠습니다.", size=10.5)

doc.add_page_break()

# ============================================================ PART 6 — 검토의견 (내부용)
h2("PART 6. [내부 검토용 — 고객 발송본에서는 반드시 제외] 워크벤치 검토의견", color=CORAL)

h3("1. 문항별 코멘트 및 추가 제안 근거")
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.LEFT
hdr = table.rows[0].cells
hdr[0].text = "문항"
hdr[1].text = "코멘트"
hdr[2].text = "관련 신규 문항"
for c in hdr:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)

review_row(table, "Q1 (ERP 활용도)",
           "원안 그대로 유효합니다. 다만 개방형 질문이었던 것을 5점 척도로 구체화했습니다 — 개방형으로 두면 응답률과 집계 편의성이 떨어집니다.",
           "-")
review_row(table, "Q2 (AI 관심도)",
           "원안 그대로 유효합니다. 5점 척도로 구체화했습니다.",
           "Q2-1 (의사결정 관여도)")
review_row(table, "Q2-1 [신규]",
           "이 문항 없이는 응답자가 실제 예산 결정권이 있는지 알 수 없어, Q5-1·Q7 응답의 신뢰도를 판단하기 어렵습니다. 리드 스코어링에 필수적입니다.",
           "-")
review_row(table, "Q3 (도입 시기)",
           "원안은 개방형 질문이었으나, \"2026~2029년 이후\" 구간을 나눠 사용자가 제시한 \"2028년 말까지\" 가설을 직접 검증할 수 있게 했습니다.",
           "-")
review_row(table, "Q4 (도입 분야)",
           "원안 그대로 유효하나 단일 선택이면 정보량이 줄어들어 복수 응답으로 전환했습니다.",
           "-")
review_row(table, "Q5 (기대 효과)",
           "복수 응답을 \"최대 2개\"로 제한해 우선순위를 드러내도록 했습니다. 전부 선택 가능하게 두면 변별력이 사라집니다.",
           "Q5-1 (예산 편성), Q5-2 (경쟁 상황)")
review_row(table, "Q5-1 [신규]",
           "\"관심\"과 \"투자 실행\"은 다른 차원입니다. 이 문항이 실질적 투자의지를 가장 직접적으로 드러내므로 반드시 포함을 권장합니다.",
           "-")
review_row(table, "Q5-2 [신규]",
           "경쟁 벤더 검토 여부를 알면 후속 컨택의 긴급도를 판단할 수 있습니다.",
           "-")
review_row(table, "Q6 (정부지원 참여의향)",
           "중요: 원안의 \"전체 사업금액의 50%, 최대 2억원 지원\"은 정부 공식 규정이 아니라 2026년 선정 4건(저스템·알피에스·선영코리아·스피폭스)을 관찰한 결과입니다(출처: 분석_영림원기존고객사_이익분석.md §2.3). 또한 \"진행 중인 사업\"이 아니라 2027년도 공고가 아직 발표되지 않은 상태입니다(출처: 05_공모전/00_공통_지원자료/공고스캔_S급품목_2026-08-04.md). 고객에게 발송되는 공식 문서에 확정 규정처럼 표기하면 이후 실제 공고 조건과 달라질 경우 신뢰도 문제가 생길 수 있어, \"2026년 선정 사례 기준\" 관찰 라벨을 반드시 유지할 것을 권장합니다.",
           "Q6-1 (과거 신청 경험)")
review_row(table, "Q6-1 [신규]",
           "과거 신청 경험 유무에 따라 \"AI사업부가 가이드\"(Q7)의 체감 가치가 크게 달라집니다 — 경험 없는 고객일수록 가이드의 가치를 높게 평가할 가능성이 있습니다.",
           "-")
review_row(table, "Q7 (공동 신청 의향)",
           "원안 그대로 유효하며, 이 설문 전체의 핵심 결론 문항입니다.",
           "Q7-1 (후속 미팅 시기)")
review_row(table, "Q7-1 [신규]",
           "이 설문의 목적이 \"의지 확인 후 실행 전환\"이므로, 의향을 실제 일정으로 전환하는 이 문항이 없으면 설문 이후 후속 조치가 지연될 위험이 큽니다. 리드 전환의 핵심 문항입니다.",
           "-")

h3("2. 추가로 고려할 만한 문항 (이번 초안에는 미포함)")
extra_items = [
    "귀사 담당자가 AI/데이터 관련 전담 인력을 보유하고 있는지 (내부 실행 역량 파악) — 담당 인력 유무에 따라 "
    "\"AI사업부 가이드\"의 필요성이 달라지므로 유용하지만, 문항 수가 이미 12개에 달해 응답 피로도를 고려해 "
    "이번 초안에서는 제외했습니다. 응답률이 낮게 나올 경우 다음 라운드에서 추가하는 것을 제안합니다.",
    "AI 도입 시 우려되는 점(데이터 보안, 초기 비용, 직원 저항 등) — 마케팅 메시지 설계에 유용하나 이 설문의 "
    "핵심 목적(도입의지·투자의지·참여가능성 확인)과는 결이 달라 별도 후속 설문으로 분리하는 것을 제안합니다.",
]
for it in extra_items:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(it)
    r.font.size = Pt(9.8)

h3("3. 설문 결과 활용 방안 제안")
para(
    "Q2(관심도)·Q5-1(예산편성)·Q7(공동신청 의향) 세 문항을 조합하면 \"도입의지 × 투자의지 × 참여가능성\" "
    "3축 리드 스코어링이 가능합니다. 예: 세 문항 모두 상위 응답(①·②)인 고객을 \"즉시 컨택 대상\"으로, "
    "일부만 해당하는 고객을 \"관찰 대상\"으로 구분해 Q7-1 응답 시기에 맞춰 방문 순서를 정하는 방식을 제안합니다. "
    "이는 기존 기획 문서의 \"고객 방문·요구사항 파악\" 단계로 자연스럽게 이어집니다.",
    size=9.8,
)

h3("4. 기존 실행 절차 기획과의 정합성")
para(
    "기존 기획(영림원AI연합군_사업전략_병합_PPTX_빌드.py NEW-B)의 3문항 가상 시나리오(\"AI 필요성 느끼는가 — "
    "거의 전원 Yes\", \"언제까지 원하는가 — 50% 2027년·50% 2029년 이내\", \"참여 의향 — 70% 이상 Yes\")는 "
    "이번 설문의 Q2, Q3, Q6에 각각 대응합니다. 실제 응답을 받으면 이 가상 응답률을 실측치로 교체해 이후 전략 "
    "자료(사업전략 병합본 등)를 갱신할 것을 권장합니다 — 해당 문서 자체에도 \"실제 확정되면 라벨 갱신 필요\"라고 "
    "명시되어 있습니다.",
    size=9.8,
)

h3("5. 발송 전 최종 확인 필요 사항")
final_checks = [
    "오르카아이티 고객 모집단 규모·특성에 대한 자료가 이 워크스페이스에 없어 서한 배경 설명에 반영하지 "
    "못했습니다 — 오르카아이티 측 확인 후 PART 1·PART 2에 반영을 권장합니다.",
    "PART 2 서한의 [ ] 표시(회사명·담당자·연락처)는 실제 발송 전 각 사에서 확정이 필요합니다.",
    "이 설문 자체는 초안이며, CLAUDE.md 완주 원칙상 실제 외부 발송(우편·이메일 등)은 항상 사람이 최종 "
    "실행합니다 — 이 문서는 발송 준비를 위한 초안입니다.",
]
for it in final_checks:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(it)
    r.font.size = Pt(9.8)

doc.save(OUT)
print("saved:", OUT)
